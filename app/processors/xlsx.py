from pathlib import Path

from openpyxl import load_workbook

from app.processors.base import BaseDocumentProcessor
from app.processors.types import DocumentElement, ParsedDocument


class XLSXProcessor(BaseDocumentProcessor):

    def can_process(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in {".xlsx", ".xlsm"}

    def process(self, file_path: Path) -> ParsedDocument:
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        elements = []

        for sheet in workbook.worksheets:
            rows = []

            for row in sheet.iter_rows(values_only=True):
                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                if any(value.strip() for value in values):
                    rows.append(values)

            if not rows:
                continue

            headers = rows[0]

            table_lines = [
                " | ".join(headers),
                " | ".join(["---"] * len(headers)),
            ]

            for row in rows[1:]:
                table_lines.append(" | ".join(row))

            table_content = "\n".join(table_lines)

            elements.append(
                DocumentElement(
                    content=table_content,
                    content_type="table",
                    metadata={
                        "sheet": sheet.title,
                    },
                )
            )

        return ParsedDocument(
            elements=elements,
            metadata={
                "file_type": "xlsx",
                "sheet_count": len(workbook.sheetnames),
                "sheets": workbook.sheetnames,
            },
        )