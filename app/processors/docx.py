from pathlib import Path

from docx import Document

from app.processors.base import BaseDocumentProcessor
from app.processors.types import DocumentElement, ParsedDocument


class DOCXProcessor(BaseDocumentProcessor):

    def can_process(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".docx"

    def process(self, file_path: Path) -> ParsedDocument:
        document = Document(file_path)

        elements = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if not text:
                continue

            content_type = (
                "heading"
                if paragraph.style.name.startswith("Heading")
                else "text"
            )

            elements.append(
                DocumentElement(
                    content=text,
                    content_type=content_type,
                    metadata={},
                )
            )

        return ParsedDocument(
            elements=elements,
            metadata={
                "file_type": "docx",
            },
        )