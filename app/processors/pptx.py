from pathlib import Path

from pptx import Presentation

from app.processors.base import BaseDocumentProcessor
from app.processors.types import DocumentElement, ParsedDocument


class PPTXProcessor(BaseDocumentProcessor):

    def can_process(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".pptx"

    def process(self, file_path: Path) -> ParsedDocument:
        presentation = Presentation(file_path)

        elements = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            texts = []

            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()

                    if text:
                        texts.append(text)

            if not texts:
                continue

            elements.append(
                DocumentElement(
                    content="\n".join(texts),
                    content_type="slide",
                    metadata={
                        "slide": slide_number,
                    },
                )
            )

        return ParsedDocument(
            elements=elements,
            metadata={
                "file_type": "pptx",
                "slide_count": len(presentation.slides),
            },
        )